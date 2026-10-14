import json, sys, re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

request = json.load(open(sys.argv[1], encoding='utf-8-sig'))
doc = Document()
for border in doc.styles.element.xpath('.//w:pBdr'):
    border.getparent().remove(border)
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(.85)
section.left_margin = section.right_margin = Inches(.95)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3']:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(8)
doc.styles['Normal'].paragraph_format.line_spacing = 1.15
doc.styles['Title'].font.size = Pt(24)
doc.styles['Title'].paragraph_format.space_after = Pt(18)
title_paragraph = doc.add_paragraph(request['title'], 'Title')
for title_run in title_paragraph.runs:
    title_run.font.name = 'Calibri'
    title_run.font.size = Pt(24)
    title_run.font.color.rgb = RGBColor(0, 0, 0)
    title_run.font.hidden = False
for block in re.split(r'\n\s*\n', request['content'].strip()):
    lines = block.splitlines()
    for line in lines:
        line = line.strip()
        if not line: continue
        if line.startswith('# ') and line[2:].strip().casefold() == request['title'].strip().casefold(): continue
        elif line.startswith('# '): doc.add_heading(line[2:], 1)
        elif line.startswith('## '): doc.add_heading(line[3:], 2)
        elif line.startswith('### '): doc.add_heading(line[4:], 3)
        elif re.match(r'^[-*] ', line): doc.add_paragraph(line[2:], 'List Bullet')
        elif re.match(r'^\d+\. ', line): doc.add_paragraph(re.sub(r'^\d+\. ', '', line), 'List Number')
        else:
            p = doc.add_paragraph()
            for i, piece in enumerate(re.split(r'\*\*(.*?)\*\*', line)):
                run = p.add_run(piece); run.bold = i % 2 == 1
footer = section.footer.paragraphs[0]
footer.alignment = 2
run = footer.add_run()
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); run._r.addnext(field)
doc.save(sys.argv[2])
print(json.dumps({'ok': True, 'paragraphs': len(doc.paragraphs)}))
